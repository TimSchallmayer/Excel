from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import main
from models import ChartSpecification


class ExcelChartIntegrationTests(unittest.TestCase):
	def _make_workbook(
		self,
		folder: str,
		headers: list[str],
		rows: list[tuple[object, ...]],
	) -> Path:
		path = Path(folder) / "messwerte.xlsx"
		workbook = Workbook()
		worksheet = workbook.active
		worksheet.title = "Messwerte"
		worksheet.append(headers)
		for row in rows:
			worksheet.append(row)
		workbook.save(path)
		return path

	def _run_and_check_chart(
		self,
		path: Path,
		x_column: str,
		y_column: str,
		fit: str = "none",
		y_error_column: str | None = None,
		origin: bool = False,
		output: Path | None = None,
		chart_type: str = "scatter",
		connect_points: bool = False,
		x_unit: str = "",
		y_unit: str = "",
		polynomial_degree: int = 2,
		moving_average_period: int = 3,
	) -> Path:
		spec = ChartSpecification(
			x_column=x_column,
			y_column=y_column,
			x_label=x_column,
			y_label=y_column,
			x_unit=x_unit,
			y_unit=y_unit,
			chart_type=chart_type,
			connect_points=connect_points,
			trendline=fit,
			polynomial_degree=polynomial_degree,
			moving_average_period=moving_average_period,
			y_error_column=y_error_column,
			origin=origin,
			confidence=0.95,
			reason="AI test decision.",
		)
		args = [str(path)]
		result_path = output or path.with_name(f"{path.stem}_PLVS_ULTRA_Graphs{path.suffix.lower()}")
		if result_path.exists():
			base = result_path
			index = 2
			while result_path.exists():
				result_path = base.with_name(f"{base.stem}_{index}{base.suffix}")
				index += 1
		if output is not None:
			args.extend(["--output", str(output)])
		with (
			patch("main.load_config", return_value={"endpoint": "https://example.invalid", "api_key": "", "model": "test"}),
			patch("main.analyze_with_ai", return_value=spec),
		):
			result = main.main(args)
		self.assertEqual(result, 0)
		self.assertTrue(result_path.is_file())
		workbook = load_workbook(result_path)
		worksheet = workbook["Messwerte"]
		self.assertEqual(worksheet["A1"].value, x_column)
		self.assertEqual(len(worksheet._charts), 1)
		chart = worksheet._charts[0]
		self.assertEqual(
			chart.__class__.__name__,
			"LineChart"
			if chart_type in {"line", "line_scatter"}
			else "BarChart"
			if chart_type in {"bar", "column"}
			else "ScatterChart",
		)
		if chart_type == "bar":
			self.assertEqual(chart.x_axis.title.tx.rich.p[0].r[0].t, y_column)
			self.assertEqual(chart.y_axis.title.tx.rich.p[0].r[0].t, x_column)
		else:
			self.assertEqual(chart.x_axis.title.tx.rich.p[0].r[0].t, x_column)
			self.assertEqual(chart.y_axis.title.tx.rich.p[0].r[0].t, y_column)
		self.assertIsNone(chart.legend)
		workbook.close()
		return result_path

	def _numeric_values(self, path: Path, column: str) -> list[float]:
		workbook = load_workbook(path, read_only=True, data_only=True)
		worksheet = workbook["Messwerte"]
		column_index = next(
			index
			for index, cell in enumerate(worksheet[1], start=1)
			if cell.value == column
		)
		values = [
			cell[0].value
			for cell in worksheet.iter_rows(
				min_row=2,
				min_col=column_index,
				max_col=column_index,
			)
		]
		workbook.close()
		return [value for value in values if isinstance(value, (int, float))]

	def test_general_numeric_table_creates_xy_chart(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = self._make_workbook(folder, ["Week", "Active users"], [(0, 1), (1, 2), (2, 4)])
			self._run_and_check_chart(path, "Week", "Active users")

	def test_axis_units_are_preserved_without_being_duplicated(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = self._make_workbook(
				folder,
				["Zeit [s]", "Strecke [m]"],
				[(0, 0), (1, 2), (2, 4)],
			)
			result = self._run_and_check_chart(
				path,
				"Zeit [s]",
				"Strecke [m]",
				x_unit="s",
				y_unit="m",
			)
			workbook = load_workbook(result)
			chart = workbook["Messwerte"]._charts[0]
			self.assertEqual(chart.x_axis.title.tx.rich.p[0].r[0].t, "Zeit [s]")
			self.assertEqual(chart.y_axis.title.tx.rich.p[0].r[0].t, "Strecke [m]")
			workbook.close()

	def test_product_categories_create_bar_and_column_charts(self) -> None:
		for chart_type, expected_type in (("bar", "bar"), ("column", "col")):
			with self.subTest(chart_type=chart_type), tempfile.TemporaryDirectory() as folder:
				path = self._make_workbook(folder, ["Produkt", "Verkäufe"], [("A", 120), ("B", 90)])
				result = self._run_and_check_chart(
					path,
					"Produkt",
					"Verkäufe",
					chart_type=chart_type,
				)
				workbook = load_workbook(result)
				chart = workbook["Messwerte"]._charts[0]
				self.assertEqual(chart.type, expected_type)
				category_reference = chart.series[0].cat.strRef or chart.series[0].cat.numRef
				self.assertEqual(category_reference.f, "'Messwerte'!$A$2:$A$3")
				workbook.close()

	def test_categorical_time_series_uses_a_line_chart_and_trendline(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = self._make_workbook(
				folder,
				["Monat", "Preis"],
				[("Jan", 12), ("Feb", 13), ("Mär", 14), ("Apr", 15), ("Mai", 17)],
			)
			result = self._run_and_check_chart(
				path,
				"Monat",
				"Preis",
				fit="linear",
				chart_type="line",
			)
			workbook = load_workbook(result)
			chart = workbook["Messwerte"]._charts[0]
			category_reference = chart.series[0].cat.strRef or chart.series[0].cat.numRef
			self.assertEqual(category_reference.f, "'Messwerte'!$A$2:$A$6")
			self.assertEqual(chart.series[0].trendline.trendlineType, "linear")
			workbook.close()

	def test_connect_points_does_not_require_a_trendline(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = self._make_workbook(folder, ["Input", "Output"], [(1, 1), (2, 2), (3, 4)])
			result = self._run_and_check_chart(
				path,
				"Input",
				"Output",
				chart_type="scatter",
				connect_points=True,
			)
			workbook = load_workbook(result)
			chart = workbook["Messwerte"]._charts[0]
			self.assertIsNone(chart.series[0].trendline)
			self.assertEqual(chart.scatterStyle, "lineMarker")
			workbook.close()

	def test_numeric_line_chart_remains_an_xy_line(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = self._make_workbook(folder, ["Input", "Output"], [(1, 1), (2, 2), (3, 4)])
			result = self._run_and_check_chart(
				path,
				"Input",
				"Output",
				chart_type="line",
			)
			workbook = load_workbook(result)
			chart = workbook["Messwerte"]._charts[0]
			self.assertEqual(chart.__class__.__name__, "LineChart")
			self.assertIsNone(chart.series[0].marker.symbol)
			self.assertIsNone(chart.series[0].trendline)
			workbook.close()

	def test_non_linear_trendlines_are_written_as_native_excel_fits(self) -> None:
		cases = (
			("exponential", [(0, 1), (1, 2), (2, 4), (3, 8)], "exp"),
			("logarithmic", [(1, -1), (2, 0), (3, 1)], "log"),
			("power", [(1, 1), (2, 4), (3, 9)], "power"),
		)
		for fit, rows, expected_type in cases:
			with self.subTest(fit=fit), tempfile.TemporaryDirectory() as folder:
				path = self._make_workbook(folder, ["Input", "Output"], rows)
				result = self._run_and_check_chart(path, "Input", "Output", fit=fit)
				workbook = load_workbook(result)
				self.assertEqual(workbook["Messwerte"]._charts[0].series[0].trendline.trendlineType, expected_type)
				workbook.close()

	def test_polynomial_degree_and_moving_average_period_are_written(self) -> None:
		rows = [(x, x * x + 1) for x in range(1, 9)]
		for fit in ("polynomial", "moving_average"):
			with self.subTest(fit=fit), tempfile.TemporaryDirectory() as folder:
				path = self._make_workbook(folder, ["Input", "Output"], rows)
				result = self._run_and_check_chart(
					path,
					"Input",
					"Output",
					fit=fit,
					polynomial_degree=4,
					moving_average_period=5,
				)
				workbook = load_workbook(result)
				trendline = workbook["Messwerte"]._charts[0].series[0].trendline
				if fit == "polynomial":
					self.assertEqual(trendline.trendlineType, "poly")
					self.assertEqual(trendline.order, 4)
				else:
					self.assertEqual(trendline.trendlineType, "movingAvg")
					self.assertEqual(trendline.period, 5)
				workbook.close()

	def test_error_column_is_added_as_native_error_bars(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = self._make_workbook(
				folder,
				["Week", "Active users", "User uncertainty"],
				[(0, 1, 0.1), (1, 2, 0.2), (2, 4, 0.1)],
			)
			result = self._run_and_check_chart(path, "Week", "Active users", y_error_column="User uncertainty")
			workbook = load_workbook(result)
			chart = workbook["Messwerte"]._charts[0]
			self.assertEqual(chart.series[0].errBars.errDir, "y")
			self.assertEqual(chart.series[0].errBars.errValType, "cust")
			workbook.close()

	def test_ai_selected_zero_origin_is_applied(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = self._make_workbook(folder, ["Input", "Output"], [(1, 1), (2, 2), (3, 4)])
			result = self._run_and_check_chart(path, "Input", "Output", origin=True)
			workbook = load_workbook(result)
			chart = workbook["Messwerte"]._charts[0]
			self.assertEqual(chart.x_axis.scaling.min, 0)
			workbook.close()

	def test_existing_output_gets_a_unique_name_and_source_is_preserved(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = self._make_workbook(folder, ["Week", "Active users"], [(0, 1), (1, 2), (2, 4)])
			original = path.read_bytes()
			first = self._run_and_check_chart(path, "Week", "Active users")
			second = self._run_and_check_chart(path, "Week", "Active users")
			self.assertNotEqual(first, second)
			self.assertEqual(path.read_bytes(), original)

	def test_existing_explicit_output_gets_a_unique_name(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = self._make_workbook(folder, ["Week", "Active users"], [(0, 1), (1, 2), (2, 4)])
			requested_output = Path(folder) / "Ergebnis.xlsx"
			requested_output.write_bytes(b"keep existing")
			result = self._run_and_check_chart(path, "Week", "Active users", output=requested_output)
			self.assertNotEqual(result, requested_output)
			self.assertEqual(requested_output.read_bytes(), b"keep existing")

	def test_empty_measurements_are_rejected(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = self._make_workbook(folder, ["Input", "Output"], [])
			self.assertEqual(main.main([str(path)]), 2)
			self.assertFalse(path.with_name(f"{path.stem}_PLVS_ULTRA_Graphs{path.suffix.lower()}").exists())

	def test_table_without_a_numeric_measurement_is_rejected(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = self._make_workbook(folder, ["Input", "Note"], [("a", "x"), ("b", "y"), ("c", "z")])
			self.assertEqual(main.main([str(path)]), 2)
			self.assertFalse(path.with_name(f"{path.stem}_PLVS_ULTRA_Graphs{path.suffix.lower()}").exists())

	def test_xlsm_is_saved_as_xlsm(self) -> None:
		with tempfile.TemporaryDirectory() as folder:
			path = Path(folder) / "messwerte.xlsm"
			workbook = Workbook()
			worksheet = workbook.active
			worksheet.title = "Data"
			worksheet.append(["Input", "Output"])
			for row in [(0, 1), (1, 2), (2, 4)]:
				worksheet.append(row)
			workbook.save(path)
			with (
				patch("main.load_config", return_value={"endpoint": "https://example.invalid", "api_key": "", "model": "test"}),
				patch("main.analyze_with_ai", return_value=ChartSpecification(
					x_column="Input",
					y_column="Output",
					x_label="Input",
					y_label="Output",
					confidence=0.95,
				)),
			):
				self.assertEqual(main.main([str(path)]), 0)
			output = path.with_name(f"{path.stem}_PLVS_ULTRA_Graphs{path.suffix.lower()}")
			self.assertTrue(output.is_file())
			result = load_workbook(output, keep_vba=True)
			self.assertEqual(len(result["Data"]._charts), 1)
			vba_archive = result.vba_archive
			result.close()
			if vba_archive is not None:
				vba_archive.close()


if __name__ == "__main__":
	unittest.main()
