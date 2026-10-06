"""Einlesen und grundlegende Prüfung von Excel-Messwerttabellen."""

from __future__ import annotations

import math
from pathlib import Path
from typing import Literal

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook as load_openpyxl_workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.series import Series as ChartSeries
from openpyxl.chart.trendline import Trendline
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from models import ChartSpecification


SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}


def load_workbook(path: Path, sheet: int | str = 0) -> pd.DataFrame:
	if path.suffix.lower() not in SUPPORTED_SUFFIXES:
		supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
		raise ValueError(f"Nicht unterstütztes Excel-Format {path.suffix!r}. Erlaubt: {supported}.")
	frame = pd.read_excel(path, sheet_name=sheet)
	column_names: list[str] = []
	counts: dict[str, int] = {}
	for column in frame.columns:
		name = str(column).strip() or "Spalte"
		counts[name] = counts.get(name, 0) + 1
		column_names.append(name if counts[name] == 1 else f"{name} ({counts[name]})")
	frame.columns = column_names
	return frame


def get_numeric_columns(frame: pd.DataFrame) -> list[str]:
	columns: list[str] = []
	for column in frame.columns:
		values = pd.to_numeric(frame[column], errors="coerce").to_numpy(dtype=float)
		finite_values = values[np.isfinite(values)]
		if finite_values.size >= 2 and np.unique(finite_values).size > 1:
			columns.append(str(column))
	return columns


def _axis_label(label: str, unit: str) -> str:
	if unit and f"[{unit}]" not in label:
		return f"{label} [{unit}]"
	return label


def _axis_bounds(
	values: np.ndarray,
	minimum: float | None,
	maximum: float | None,
	use_zero: bool,
) -> tuple[float, float]:
	data_min = float(np.min(values))
	data_max = float(np.max(values))
	span = data_max - data_min
	padding = span * 0.08 if span else max(abs(data_min) * 0.08, 1.0)
	low = minimum if minimum is not None else data_min - padding
	high = maximum if maximum is not None else data_max + padding
	if use_zero and data_min >= 0 and (minimum is None or minimum <= 0):
		low = 0.0
	if use_zero and data_max <= 0 and (maximum is None or maximum >= 0):
		high = 0.0
	if not math.isfinite(low) or not math.isfinite(high) or low >= high:
		raise ValueError("Die Diagramm-Achsengrenzen sind ungültig.")
	if data_min < 0 < data_max and minimum is None and maximum is None and not use_zero:
		low = min(low, 0.0)
		if high > 0:
			high = max(high, 0.0)
	return low, high


def _use_zero_origin(origin: bool | None, values: np.ndarray) -> bool:
	if origin is True:
		return float(np.min(values)) >= 0
	return False


def _append_error_bars(
	series: ChartSeries,
	worksheet: Worksheet,
	column: int,
	first_row: int,
	last_row: int,
	direction: Literal["x", "y"],
) -> None:
	column_letter = get_column_letter(column)
	reference = NumDataSource(
		numRef=NumRef(f=f"'{worksheet.title}'!${column_letter}${first_row}:${column_letter}${last_row}")
	)
	series.errBars = ErrorBars(
		errDir=direction,
		errValType="cust",
		plus=reference,
		minus=reference,
	)


def create_excel_chart(
	workbook: Workbook,
	data_sheet: Worksheet,
	frame: pd.DataFrame,
	spec: ChartSpecification,
	fit: str = "none",
) -> BarChart | LineChart | ScatterChart:
	"""Fügt ein natives Excel-Diagramm rechts neben den Messdaten ein."""
	if fit not in {"none", "linear", "quadratic", "cubic", "polynomial", "exponential", "logarithmic", "power", "moving_average"}:
		raise ValueError(f"Nicht unterstützte Trendlinie: {fit}")
	if spec.chart_type not in {"scatter", "scatter_lines", "line", "line_scatter", "bar", "column"}:
		raise ValueError(f"Nicht unterstützte Diagrammart: {spec.chart_type}")
	for column in (spec.x_column, spec.y_column):
		if column not in frame.columns:
			raise ValueError(f"Diagrammspalte {column!r} fehlt in der Tabelle.")

	x_values = pd.to_numeric(frame[spec.x_column], errors="coerce").to_numpy(dtype=float)
	y_values = pd.to_numeric(frame[spec.y_column], errors="coerce").to_numpy(dtype=float)
	y_valid = np.isfinite(y_values)
	categorical_chart = spec.chart_type in {"bar", "column", "line", "line_scatter"} or not np.any(np.isfinite(x_values))
	if categorical_chart:
		if spec.chart_type in {"scatter", "scatter_lines"}:
			raise ValueError("Ein Scatter-/XY-Diagramm benötigt eine numerische x-Spalte.")
		x_categories = frame[spec.x_column]
		valid = x_categories.notna().to_numpy() & y_valid
		if int(valid.sum()) < 2:
			raise ValueError("Für ein Diagramm werden mindestens zwei gültige Datenpaare benötigt.")
		category_codes = pd.factorize(x_categories, sort=False)[0].astype(float) + 1
		x_data = category_codes[valid]
	else:
		valid = np.isfinite(x_values) & y_valid
		if valid.sum() < 2 or np.unique(x_values[valid]).size < 2:
			raise ValueError("Für ein XY-Diagramm werden mindestens zwei gültige Messwertpaare mit unterschiedlichen x-Werten benötigt.")
		x_data = x_values[valid]
	y_data = y_values[valid]
	if fit != "none":
		if fit in {"linear", "quadratic", "cubic", "polynomial"}:
			degree = (
				1 if fit == "linear"
				else 2 if fit == "quadratic"
				else 3 if fit == "cubic"
				else spec.polynomial_degree
			)
			if fit == "polynomial" and not 2 <= degree <= 6:
				raise ValueError("Der Polynomgrad muss zwischen 2 und 6 liegen.")
			if np.unique(x_data).size <= degree:
				raise ValueError(f"Ein Fit vom Grad {degree} benötigt mindestens {degree + 1} unterschiedliche x-Werte.")
		if fit == "exponential":
			if np.any(y_data <= 0):
				raise ValueError("Ein exponentieller Fit erfordert positive y-Werte.")
		if fit in {"logarithmic", "power"} and np.any(x_data <= 0):
			raise ValueError("Für logarithmische und Potenz-Fits müssen x-Werte positiv sein.")
		if fit == "power" and np.any(y_data <= 0):
			raise ValueError("Für einen Potenz-Fit müssen y-Werte positiv sein.")
		if fit == "moving_average":
			if spec.moving_average_period not in {2, 3, 4, 5, 6, 7, 10, 12, 20}:
				raise ValueError("Die Periode des gleitenden Durchschnitts ist ungültig.")
			if spec.moving_average_period >= len(y_data):
				raise ValueError(
					"Der gleitende Durchschnitt benötigt mehr Messpunkte als die gewählte Periode."
				)

	x_column_index = list(frame.columns).index(spec.x_column) + 1
	y_column_index = list(frame.columns).index(spec.y_column) + 1
	x_reference = Reference(data_sheet, min_col=x_column_index, min_row=2, max_row=data_sheet.max_row)
	y_reference = Reference(data_sheet, min_col=y_column_index, min_row=2, max_row=data_sheet.max_row)
	if spec.chart_type in {"bar", "column"}:
		chart = BarChart()
		chart.type = "bar" if spec.chart_type == "bar" else "col"
		chart.grouping = "clustered"
		chart.add_data(y_reference, titles_from_data=False)
		chart.set_categories(x_reference)
		series = chart.series[0]
	elif spec.chart_type in {"line", "line_scatter"} or categorical_chart:
		chart = LineChart()
		chart.add_data(y_reference, titles_from_data=False)
		chart.set_categories(x_reference)
		series = chart.series[0]
	else:
		chart = ScatterChart()
		connect = spec.connect_points or spec.chart_type == "scatter_lines"
		chart.scatterStyle = (
			"lineMarker" if connect and spec.show_points
			else "line" if connect
			else "marker"
		)
		series = Series(y_reference, x_reference, title="Messwerte")
	chart.title = f"{spec.y_label} in Abhängigkeit von {spec.x_label}"
	if isinstance(chart, BarChart) and chart.type == "bar":
		chart.x_axis.title = _axis_label(spec.y_label, spec.y_unit)
		chart.y_axis.title = _axis_label(spec.x_label, spec.x_unit)
	else:
		chart.x_axis.title = _axis_label(spec.x_label, spec.x_unit)
		chart.y_axis.title = _axis_label(spec.y_label, spec.y_unit)
	chart.style = 13
	chart.width = 18
	chart.height = 10
	chart.legend = None
	chart.display_blanks = "gap"
	chart.visible_cells_only = False

	if spec.chart_type in {"scatter", "scatter_lines", "line_scatter"} and spec.show_points:
		series.marker.symbol = "circle"
		series.marker.size = 7
	elif spec.chart_type in {"line", "scatter", "scatter_lines", "line_scatter"}:
		series.marker.symbol = "none"
	if isinstance(chart, ScatterChart) and not (
		spec.connect_points or spec.chart_type == "scatter_lines"
	):
		series.graphicalProperties.line.noFill = True
	if fit != "none":
		trendline_type = {
			"linear": "linear",
			"quadratic": "poly",
			"cubic": "poly",
			"polynomial": "poly",
			"exponential": "exp",
			"logarithmic": "log",
			"power": "power",
			"moving_average": "movingAvg",
		}[fit]
		trendline_kwargs = {
			"name": f"{fit.title()} Fit",
			"trendlineType": trendline_type,
			"dispRSqr": True,
		}
		if fit in {"quadratic", "cubic", "polynomial"}:
			trendline_kwargs["order"] = (
				2 if fit == "quadratic"
				else 3 if fit == "cubic"
				else spec.polynomial_degree
			)
		elif fit == "moving_average":
			trendline_kwargs["period"] = spec.moving_average_period
		series.trendline = Trendline(**trendline_kwargs)
	if isinstance(chart, ScatterChart):
		chart.series.append(series)

	error_bar_columns: list[tuple[int, Literal["x", "y"]]] = []
	for error_column, direction in (
		(spec.x_error_column, "x"),
		(spec.y_error_column, "y"),
	):
		if not error_column:
			continue
		if error_column not in frame.columns:
			raise ValueError(f"Fehlerwert-Spalte {error_column!r} existiert nicht.")
		error_column_index = list(frame.columns).index(error_column) + 1
		error_values = pd.to_numeric(frame[error_column], errors="coerce").to_numpy(dtype=float)[valid]
		if not np.any(np.isfinite(error_values) & (error_values >= 0)):
			continue
		error_bar_columns.append((error_column_index, direction))
	if error_bar_columns:
		preferred_direction = "y" if any(direction == "y" for _, direction in error_bar_columns) else "x"
		selected_error_bars = next(
			(error for error in error_bar_columns if error[1] == preferred_direction),
			error_bar_columns[0],
		)
		_append_error_bars(series, data_sheet, selected_error_bars[0], 2, data_sheet.max_row, selected_error_bars[1])

	y_zero = _use_zero_origin(spec.origin, values=y_data)
	y_min, y_max = _axis_bounds(y_data, spec.y_min, spec.y_max, y_zero)
	if isinstance(chart, BarChart) and chart.type == "bar":
		chart.x_axis.scaling.min = y_min
		chart.x_axis.scaling.max = y_max
	elif isinstance(chart, BarChart):
		chart.y_axis.scaling.min = y_min
		chart.y_axis.scaling.max = y_max
	else:
		chart.y_axis.scaling.min = y_min
		chart.y_axis.scaling.max = y_max
		if not categorical_chart:
			x_zero = _use_zero_origin(spec.origin, values=x_data)
			x_min, x_max = _axis_bounds(x_data, spec.x_min, spec.x_max, x_zero)
			chart.x_axis.scaling.min = x_min
			chart.x_axis.scaling.max = x_max

	data_sheet.add_chart(chart, f"{get_column_letter(data_sheet.max_column + 2)}2")
	return chart


def output_path_for(input_path: Path) -> Path:
	"""Wählt einen neuen Namen, ohne vorhandene Dateien zu überschreiben."""
	base = input_path.with_name(f"{input_path.stem}_PLVS_ULTRA_Graphs{input_path.suffix.lower()}")
	return _non_existing_path(base)


def _non_existing_path(path: Path) -> Path:
	base = path
	candidate = base
	index = 2
	while candidate.exists():
		candidate = base.with_name(f"{base.stem}_{index}{base.suffix}")
		index += 1
	return candidate


def save_workbook_with_chart(
	input_path: Path,
	frame: pd.DataFrame,
	sheet: int | str,
	spec: ChartSpecification,
	fit: str = "none",
	output_path: Path | None = None,
) -> Path:
	"""Speichert eine Arbeitsmappenkopie mit nativem Excel-Diagramm."""
	input_path = input_path.resolve()
	if input_path.suffix.lower() not in {".xlsx", ".xlsm"}:
		raise ValueError("Es werden nur .xlsx- und .xlsm-Dateien unterstützt.")
	output = output_path or output_path_for(input_path)
	output = output.resolve()
	if output == input_path:
		raise ValueError("Die Ausgabedatei darf nicht die Eingabedatei überschreiben.")
	if output.suffix.lower() != input_path.suffix.lower():
		raise ValueError("Eingabe- und Ausgabeformat müssen übereinstimmen.")
	if output_path is not None:
		output = _non_existing_path(output)
	output.parent.mkdir(parents=True, exist_ok=True)
	workbook = load_openpyxl_workbook(input_path, keep_vba=input_path.suffix.lower() == ".xlsm")
	try:
		if isinstance(sheet, int):
			if sheet < 0 or sheet >= len(workbook.worksheets):
				raise ValueError(f"Tabellenblatt-Index {sheet} ist nicht vorhanden.")
			data_sheet = workbook.worksheets[sheet]
		else:
			if sheet not in workbook.sheetnames:
				raise ValueError(f"Tabellenblatt {sheet!r} ist nicht vorhanden.")
			data_sheet = workbook[sheet]
		create_excel_chart(workbook, data_sheet, frame, spec, fit)
		workbook.save(output)
	finally:
		workbook.close()
		vba_archive = getattr(workbook, "vba_archive", None)
		if vba_archive is not None:
			vba_archive.close()
	return output